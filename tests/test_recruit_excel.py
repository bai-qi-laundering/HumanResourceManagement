import io
import json
import unittest
from urllib.parse import urlencode

import openpyxl
import test_server
from recruitment_excel import read_excel


class RecruitmentExcelTests(unittest.TestCase):
    setUp = test_server.HRTests.setUp
    tearDown = test_server.HRTests.tearDown
    call = test_server.HRTests.call
    emp = test_server.HRTests.emp

    def source(self):
        # Synthetic workbook only; no real candidate data is committed.
        w = openpyxl.Workbook()
        w.remove(w.active)
        header = ['姓名', '電話號碼', '職位', '性別', '年齡', '面試時間', '備註', '人力網', '撥打時間', '再次撥打']
        for name in ['115年6月30日', '115年7月1日', '115年7月2日8日']:
            s = w.create_sheet(name)
            s.append(header)
            s.append(['測試人員', '0900-000-001', '作業員', '', '', None, '未接來電', '104', '10:00', ''])
        s = w.create_sheet('115年7月9日')
        s.append(header)
        s.append(['只有名單', '0900000002', '作業員'])
        s = w.create_sheet('每日電話')
        s.append(['日期', '姓名', '電話', '職務', '備註'])
        s.append(['2026/06/30', '六月測試', '0900000003', '作業員', '六月不匯入'])
        s.append(['2026/07/19', '測試人員', '0900000001', '作業員', '有興趣'])
        b = io.BytesIO()
        w.save(b)
        return b.getvalue()

    def payload(self):
        return read_excel(self.source(), 'phone.xlsx', '2026-07-01', '2026-10-20')

    def test_dates_ranges_and_no_invented_calls(self):
        p = self.payload()
        self.assertEqual(len(p['records']), 4)
        self.assertTrue(all(r['rangeStart'] >= '2026-07-01' for r in p['records']))
        multi = next(r for r in p['records'] if r['sourceSheet'] == '115年7月2日8日')
        self.assertEqual(multi['callDate'], '')
        self.assertEqual((multi['rangeStart'], multi['rangeEnd']), ('2026-07-02', '2026-07-08'))
        preview = self.call('/api/recruitment/import-preview', p)
        self.assertEqual(preview['data']['summary']['people'], 2)
        self.assertEqual(preview['data']['summary']['unknownDates'], 1)
        self.assertEqual(len(self.call('/api/state')['data']['candidates']), 0)
        result = self.call('/api/recruitment/import', p)
        self.assertEqual(result['status'], 200, result)
        self.assertEqual(result['data']['counts'], {'candidates': 2, 'rows': 4, 'events': 3, 'skipped': 0})
        state = self.call('/api/state')['data']
        unknown = next(e for e in state['recruitmentEvents'] if not e['contactDate'])
        self.assertEqual(unknown['at'], '')
        again = self.call('/api/recruitment/import', p)['data']['counts']
        self.assertEqual(again, {'candidates': 0, 'rows': 0, 'events': 0, 'skipped': 4})
        self.assertEqual(len(self.call('/api/state')['data']['recruitmentEvents']), 3)

    def test_manual_profile_and_followup_preserved(self):
        c = self.call('/api/recruitment/candidate', {'name': '測試人員', 'phone': '0900000001', 'department': '洗滌', 'status': '待報到', 'notes': '人工決定'})['data']
        self.call('/api/recruitment/import', self.payload())
        c = next(v for v in self.call('/api/state')['data']['candidates'] if v['id'] == c['id'])
        self.assertEqual(c['status'], '待報到')
        self.assertEqual(c['notes'], '人工決定')
        self.assertEqual(c['department'], '洗滌')
        invalid = self.payload()
        invalid['records'][0]['rangeStart'] = '2026-06-30'
        self.assertEqual(self.call('/api/recruitment/import', invalid)['status'], 400)

    def test_excel_endpoint_preview_and_network_guard(self):
        raw = self.source()
        env = {'REQUEST_METHOD': 'POST', 'PATH_INFO': '/api/recruitment/excel-preview', 'QUERY_STRING': urlencode({'filename': 'test.xlsx', 'start': '2026-07-01', 'end': '2026-10-20'}), 'REMOTE_ADDR': '192.168.1.50', 'HTTP_HOST': '192.168.1.20:8090', 'wsgi.url_scheme': 'http', 'wsgi.input': io.BytesIO(raw), 'CONTENT_LENGTH': str(len(raw)), 'HTTP_COOKIE': self.cookie, 'HTTP_X_HR_CSRF': self.csrf}
        result = {}
        def start(status, headers):result['status'] = int(status.split()[0])
        response = json.loads(b''.join(self.app(env, start)))
        self.assertEqual(result['status'], 200, response)
        self.assertEqual(response['summary']['rows'], 4)
        env['REMOTE_ADDR'] = '100.90.1.1'
        self.app(env, start)
        self.assertEqual(result['status'], 403)

    def test_inherited_year_rollover(self):
        w=openpyxl.Workbook();w.remove(w.active)
        for title in ['111年12月28日','1月3日']:
            sh=w.create_sheet(title);sh.append(['姓名','電話']);sh.append(['跨年測試','0900000009'])
        b=io.BytesIO();w.save(b)
        result=read_excel(b.getvalue(),'test.xlsx','2023-01-01','2023-01-31')
        self.assertEqual(len(result['records']),1)
        self.assertEqual(result['records'][0]['callDate'],'2023-01-03')
