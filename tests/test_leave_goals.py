import io
import unittest
from datetime import datetime, timedelta

import openpyxl
import test_server
from leave_excel import read_excel
from server import TZ, Store


class LeaveGoalTests(unittest.TestCase):
    setUp = test_server.HRTests.setUp
    tearDown = test_server.HRTests.tearDown
    call = test_server.HRTests.call
    emp = test_server.HRTests.emp

    def progress(self):
        return self.call('/api/state')['data']['recruitmentGoals'][0]

    def test_goals_onboarding_departure_increase_and_conflicts(self):
        today = datetime.now(TZ).date()
        yesterday = (today - timedelta(days=1)).isoformat()
        tomorrow = (today + timedelta(days=1)).isoformat()
        a = self.emp()
        self.emp('F002', '測試二')
        result = self.call('/api/recruitment/goal', {'department': '摺疊', 'additional': 5})
        self.assertEqual(result['status'], 200)
        self.assertEqual((self.progress()['baseline'], self.progress()['percent']), (2, 0))
        self.emp('F003', '新員工一', arrivalDate=today.isoformat())
        self.emp('F004', '新員工二', arrivalDate=today.isoformat())
        self.emp('F005', '未到職', arrivalDate=tomorrow)
        self.emp('W001', '其他部門', department='洗滌')
        self.assertEqual((self.progress()['netFilled'], self.progress()['percent']), (2, 40))
        self.call('/api/employee', {**a, 'departureDate': yesterday})
        self.assertEqual((self.progress()['remaining'], self.progress()['percent']), (4, 20))
        self.assertEqual(self.call('/api/recruitment/goal', {'department': '摺疊', 'additional': 2, 'revision': 1})['status'], 200)
        g = self.progress()
        self.assertEqual((g['baseline'], g['required'], g['netFilled'], len(g['history'])), (2, 7, 1, 2))
        self.assertEqual(self.call('/api/recruitment/goal', {'department': '摺疊', 'additional': 2, 'revision': 1})['status'], 409)
        self.assertEqual(self.call('/api/recruitment/goal', {'department': '洗滌', 'additional': 1.5})['status'], 400)
        fresh = Store(self.app.store.path)
        self.assertEqual(fresh.state({})['recruitmentGoals'][0]['required'], 7)

    def test_conversion_and_source_sync_affect_goal_without_counting_candidate_status(self):
        self.emp()
        self.call('/api/recruitment/goal', {'department': '摺疊', 'additional': 2})
        c = self.call('/api/recruitment/candidate', {'name': '測試報到', 'phone': '0900000001', 'department': '洗滌', 'status': '已報到'})['data']
        self.assertEqual(self.progress()['filled'], 0)
        response = self.call('/api/recruitment/convert', {'candidateId': c['id'], 'revision': c['revision'], 'employeeId': 'F100', 'department': '摺疊', 'arrivalDate': '2024-01-01', 'employmentType': '正職'})
        self.assertEqual(response['status'], 200, response)
        self.assertEqual((self.progress()['filled'], self.progress()['percent']), (1, 50))
        # A corrected departure from GAS is reflected by the same projection.
        self.call('/api/import', {'schema': 1, 'employees': [{'id': 'F001', 'name': '王小明', 'department': '摺疊', 'arrivalDate': '2024-04-01', 'departureDate': '2024-04-02'}]})
        self.assertEqual(self.progress()['filled'], 0)

    def source(self):
        w = openpyxl.Workbook(); w.remove(w.active)
        for sheet, arrival in [('114', 1120121), ('115', 1130121)]:
            sh = w.create_sheet(sheet)
            sh.append(['姓名', '到職日', sheet+'年結算日', '天數', '已休', '備註'])
            sh.append(['王小明', arrival, 1140121, '7+3', '', '已於薪資結清，歷史註記'])
            sh.append(['同名測試', 1120201, 1140201, 7, '3+1', '核對'])
        b = io.BytesIO(); w.save(b)
        return b.getvalue()

    def test_leave_preview_history_dedup_conflicts_and_manual_binding(self):
        self.emp()
        self.emp('F002', '同名測試'); self.emp('F003', '同名測試')
        p = read_excel(self.source(), 'synthetic.xlsx')
        preview = self.app.store.preview_leave(p)
        r = preview['rows'][0]
        self.assertEqual((r['arrivalDate'], r['days'], r['usedDays']), ('2023-01-21', 10, None))
        self.assertTrue(any('不一致' in w for w in r['reviewWarnings']))
        self.assertTrue(any('不同' in w for w in r['reviewWarnings']))
        self.assertEqual(preview['rows'][1]['employeeId'], '')
        self.assertEqual(len(self.call('/api/state')['data']['leaveExcel']), 0)
        self.assertEqual(self.call('/api/leave/import', p)['data'], {'added': 4, 'skipped': 0})
        state = self.call('/api/state')['data']
        self.assertEqual(state['records'], [])
        self.assertEqual(state['employees'][0]['arrivalDate'], '2024-04-01')
        row = next(r for r in state['leaveExcel'] if r['name'] == '同名測試')
        self.assertIsNone(row['employeeId'])
        self.assertEqual(self.call('/api/leave/bind', {'id': row['id'], 'employeeId': 'F002'})['status'], 200)
        p['records'][0]['sourceFile'] = 'renamed.xlsx'
        self.assertEqual(self.call('/api/leave/import', p)['data'], {'added': 0, 'skipped': 4})
        self.assertTrue(all(r['repeated'] for r in self.app.store.preview_leave(p)['rows']))
        bound = next(r for r in self.call('/api/state')['data']['leaveExcel'] if r['id'] == row['id'])
        self.assertEqual(bound['employeeId'], 'F002')
        # Mismatched settlement year is surfaced, never converted to a grant.
        self.assertTrue(any(r['warnings'] for r in state['leaveExcel'] if r['sourceSheet'] == '115'))

    def test_leave_date_aliases_and_expression_not_executed(self):
        w = openpyxl.Workbook(); sh = w.active
        sh.append(['員工', '入職日', '已休天數', '結算日', '天數'])
        sh.append(['測試員工', '115/07/01', '3＋1', '116/07/01', '__import__("os")'])
        b = io.BytesIO(); w.save(b)
        row = read_excel(b.getvalue(), 'test.xlsx')['records'][0]
        self.assertEqual((row['arrivalDate'], row['usedDays'], row['days']), ('2026-07-01', 4, None))
        self.assertTrue(row['warnings'])


if __name__ == '__main__':
    unittest.main()
